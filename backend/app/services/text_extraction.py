"""Plain-text extraction from validated PDF/DOCX uploads.

Inputs: raw file bytes plus the kind already detected by content inspection.
Uses no Gemini key — extraction is pure-Python (pypdf / python-docx) so the
LLM only ever sees text, never raw binary.
"""

import io


class TextExtractionError(ValueError):
    """Raised with a user-facing message when no usable text is found."""


_MIN_USABLE_CHARS = 50


def extract_text(data: bytes, kind: str) -> str:
    """Return the document's text or raise ``TextExtractionError``."""
    if kind == "pdf":
        text = _extract_pdf(data)
    elif kind == "docx":
        text = _extract_docx(data)
    else:
        raise TextExtractionError(f"Unsupported file kind: {kind}")

    text = text.strip()
    if len(text) < _MIN_USABLE_CHARS:
        raise TextExtractionError(
            "We couldn't read any text from this file. If it is a scanned or "
            "image-only PDF, please upload a text-based export instead."
        )
    return text


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # noqa: BLE001
        raise TextExtractionError(
            "This PDF could not be read — it may be corrupted or password-protected."
        ) from exc


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise TextExtractionError(
            "This DOCX file could not be read — it may be corrupted."
        ) from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
